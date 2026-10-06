import io
from unittest.mock import AsyncMock

import pytest
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from llm_wiki.agents.visual_presentation import export_deck, plan_deck, render_slide, validate_image


@pytest.mark.parametrize(
    "language,title",
    [
        ("ru", "Качество и доверие"),
        ("kk", "Әділдік, қауіпсіздік, ашықтық"),
        ("en", "Quality and trust"),
    ],
)
@pytest.mark.parametrize("layout", ["hero", "editorial", "comparison", "big_number"])
@pytest.mark.parametrize("size", [(1536, 864), (2048, 1152)])
def test_native_dimensions_and_exports(language, title, layout, size):
    slide = {
        "title": title,
        "subtitle": "1 500 / 1,500",
        "points": [title, "20"],
        "layout": layout,
        "source_refs": ["f"],
        "speaker_notes": "Exact notes",
    }
    data = render_slide(slide, None, 1, 1, language, size)
    validate_image(data, size)
    assert Image.open(io.BytesIO(data)).size == size
    pptx = Presentation(io.BytesIO(export_deck([data], [slide], "pptx")))
    assert pptx.slide_width * 9 == pptx.slide_height * 16
    assert pptx.slides[0].notes_slide.notes_text_frame.text == "Exact notes"
    assert pptx.slides[0].shapes[0].image.blob == data
    pdf = PdfReader(io.BytesIO(export_deck([data], [slide], "pdf")))
    assert len(pdf.pages) == 1
    assert float(pdf.pages[0].mediabox.width) / float(pdf.pages[0].mediabox.height) == 16 / 9


def test_wrong_aspect_and_overflow_fail_instead_of_cropping():
    with pytest.raises(ValueError, match="image_dimensions_invalid"):
        render_slide({}, None, 1, 1, "ru", (1536, 1024))
    with pytest.raises(ValueError, match="slide_text_overflow"):
        render_slide(
            {"title": "text " * 200, "layout": "hero", "source_refs": ["f"]}, None, 1, 1, "en"
        )


@pytest.mark.asyncio
async def test_planner_checks_provenance_before_images_and_can_change_layout(monkeypatch):
    import json

    import llm_wiki.llm.client as client

    slide = {
        "title": "Short title",
        "subtitle": "",
        "points": ["20 sites"],
        "layout": "hero",
        "source_refs": ["f"],
        "speaker_notes": "",
        "image_brief": "Glass architecture",
    }
    deck = {"title": "Deck", "style_block": "Blue glass", "warnings": [], "slides": [slide]}
    mock = AsyncMock()
    mock.complete.return_value = (json.dumps(deck), {})
    monkeypatch.setattr(client, "LLMClient", lambda: mock)
    import llm_wiki.agents.visual_presentation as visual

    original = visual.render_slide

    def rejecting_hero(slide, *args):
        if slide["layout"] == "hero":
            raise ValueError("slide_text_overflow")
        return original(slide, *args)

    monkeypatch.setattr(visual, "render_slide", rejecting_hero)
    result = await plan_deck(
        {"title": "Deck", "sources": [{"id": "f", "name": "Evidence", "text": "20 sites"}]},
        "en",
        "job",
    )
    assert result["slides"][0]["layout"] == "editorial"
    assert result["slides"][0]["points"] == ["20 sites"]
    assert result["slides"][0]["source_numbers"] == [1]
    deck["slides"][0]["source_refs"] = ["invented"]
    mock.complete.return_value = (json.dumps(deck), {})
    with pytest.raises(ValueError, match="deck_invalid"):
        await plan_deck(
            {"title": "Deck", "sources": [{"id": "f", "name": "Evidence", "text": "20 sites"}]},
            "en",
            "job",
        )
