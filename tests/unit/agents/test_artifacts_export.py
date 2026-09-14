"""Unit tests for artifact export (docx / pptx / pdf byte generation)."""

from __future__ import annotations

import io
import zipfile

import pytest

from llm_wiki.agents.artifacts_export import (
    ExportError,
    _first_existing,
    _REGULAR_FONT_CANDIDATES,
    export_artifact,
    supported_formats,
)

REPORT = {
    "summary": "Резюме материала о летнем лагере.",
    "key_insight": "Главный вывод одним предложением.",
    "risks": ["Нет владельца процесса.", "Метрики собираются вручную."],
    "recommendations": ["Пилот — запустить на 6 недель."],
    "sources": ["Страница-источник"],
    "relevance_pct": 93,
    "citation_coverage_pct": 95,
    "effect_horizon": "6–12 мес",
    "source_language": "RU",
    "reading_minutes": 3,
}
REPORT_LEGACY = {
    "executive_summary": "Резюме материала (старый формат).",
    "metrics": [{"label": "Формат", "value": "Лагерь"}],
    "sections": [{"heading": "Замысел", "body": "Тело раздела с кириллицей."}],
}
TEST = {
    "questions": [
        {"prompt": "Что это?", "options": ["Вариант А", "Вариант Б"], "correct": 1,
         "explanation": "Потому что Б."},
    ]
}
CARD = {
    "insight": "Главный вывод одним предложением.",
    "context": "Что за материал и зачем.",
    "steps": [{"title": "Шаг A", "text": "Пояснение шага."}],
    "risk": "Что может сломаться.",
    "action": "Собрать команду на 60 минут.",
    "action_minutes": 60,
    "relevance_pct": 93,
    "source_language": "RU",
}
CARD_LEGACY = {
    "title": "Карточка знаний",
    "summary": "Краткий итог.",
    "key_points": [{"label": "Инсайт", "text": "Важная мысль."}],
    "recommendations": ["Сделать шаг X"],
    "tags": ["обучение", "практика"],
}
PRESENTATION = {
    "title": "Дека о лагере",
    "slides": [{"heading": "Слайд 1", "bullets": ["раз", "два"], "notes": "заметки"}],
}

_HAS_PDF_FONT = _first_existing(_REGULAR_FONT_CANDIDATES) is not None
needs_font = pytest.mark.skipif(not _HAS_PDF_FONT, reason="no Unicode TTF on this host")


def _is_zip_with(data: bytes, member_prefix: str) -> bool:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return any(n.startswith(member_prefix) for n in z.namelist())


# --- registry -------------------------------------------------------------


def test_supported_formats() -> None:
    assert supported_formats("report") == {"pdf", "docx"}
    assert supported_formats("presentation") == {"pdf", "pptx"}
    assert supported_formats("infographic") == set()
    assert supported_formats("podcast") == set()


@pytest.mark.parametrize("kind,fmt", [
    ("report", "pptx"), ("presentation", "docx"), ("infographic", "pdf"), ("card", "xml"),
])
def test_unsupported_pair_raises(kind: str, fmt: str) -> None:
    with pytest.raises(ExportError):
        export_artifact(kind, {}, fmt)


# --- docx (OOXML zip) -----------------------------------------------------


@pytest.mark.parametrize("kind,content", [
    ("report", REPORT), ("report", REPORT_LEGACY), ("test", TEST),
    ("card", CARD), ("card", CARD_LEGACY),
])
def test_docx_is_valid_ooxml(kind: str, content: dict) -> None:
    data, media = export_artifact(kind, content, "docx")
    assert media.endswith("wordprocessingml.document")
    assert data[:2] == b"PK"  # zip magic
    assert _is_zip_with(data, "word/document.xml")


# --- pptx (OOXML zip) -----------------------------------------------------


def test_pptx_has_slides() -> None:
    data, media = export_artifact("presentation", PRESENTATION, "pptx")
    assert media.endswith("presentationml.presentation")
    assert data[:2] == b"PK"
    assert _is_zip_with(data, "ppt/slides/slide")


# --- pdf (needs a Unicode font) -------------------------------------------


@needs_font
@pytest.mark.parametrize("kind,content", [
    ("report", REPORT), ("report", REPORT_LEGACY), ("test", TEST),
    ("card", CARD), ("card", CARD_LEGACY), ("presentation", PRESENTATION),
])
def test_pdf_magic(kind: str, content: dict) -> None:
    data, media = export_artifact(kind, content, "pdf")
    assert media == "application/pdf"
    assert data[:4] == b"%PDF"
    assert len(data) > 500  # non-trivial document


@needs_font
def test_pdf_handles_empty_content() -> None:
    # missing keys must not crash — just a near-empty document
    data, _ = export_artifact("report", {}, "pdf")
    assert data[:4] == b"%PDF"


def test_pptx_matches_preview_geometry_without_extra_cover():
    from pptx import Presentation
    from pptx.util import Pt
    from llm_wiki.agents.presentation_layout import presentation_layout
    content = {"title": "Do not add a cover", "slides": [
        {"heading": "Обзор кейса", "bullets": ["Контекст и цели", "Риски и решения"], "notes": "Примечание докладчика"},
        {"title": "Ұсыныстар", "body": "Бизнес және команда"},
        {"notes": "Invalid empty slide"},
    ]}
    layout = presentation_layout(content)
    data, _ = export_artifact("presentation", content, "pptx")
    deck = Presentation(io.BytesIO(data))
    assert len(deck.slides) == len(layout["slides"]) == 2
    assert deck.slide_width == Pt(960) and deck.slide_height == Pt(540)
    assert "Примечание докладчика" in deck.slides[0].notes_slide.notes_text_frame.text
    for slide, expected in zip(deck.slides, layout["slides"]):
        texts = [shape for shape in slide.shapes if shape.has_text_frame and shape.text]
        assert [shape.text for shape in texts] == [block["text"] for block in expected["blocks"]]
        for shape, block in zip(texts, expected["blocks"]):
            assert shape.left == Pt(block["x"]) and shape.top == Pt(block["y"])
            assert shape.text_frame.paragraphs[0].font.size == Pt(block["size"])


def test_dense_slides_fit_without_losing_any_bullets():
    from llm_wiki.agents.presentation_layout import presentation_layout
    bullets = [f"{i}. " + "Очень длинное описание задачи и результата. " * 7 for i in range(12)]
    slide = presentation_layout({"slides": [{"heading": "План", "bullets": bullets}]})["slides"][0]
    assert len(slide["marks"]) == len(bullets)
    assert all(0 <= b["y"] and b["y"] + b["size"] <= 540 for b in slide["blocks"])
    assert "".join(b["text"] for b in slide["blocks"]).replace(" ", "") == ("План" + "".join(bullets)).replace(" ", "")
